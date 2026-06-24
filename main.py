from pathlib import Path
import pandas as pd
import openpyxl as xl
from openpyxl.styles import Font, PatternFill, Border, Side

#Styles
header_font = Font(bold=True)
header_fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")

border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

def auto_width(ws):
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter

        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))

        ws.column_dimensions[col_letter].width = max_length + 2


def style_table(ws):
    for row in ws.iter_rows():
        for cell in row:
            cell.border = border

    # header
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill


day1 = Path('./src/day1')
day2 = Path('./src/day2')
day3 = Path('./src/day3')


feedback_df = pd.DataFrame(
    columns=['identite', 'jour', 'score', 'marquant', 'activite', 'satisfaction', 'besoin', 'objectif', 'atout'])
transport_df = pd.DataFrame(columns=['identite', 'transport', 'duree', 'max'])
logement_df = pd.DataFrame(columns=['identite', 'logement', 'stable', 'changement',
             'justif', 'enfant'])
capacity_df = pd.DataFrame(columns=['identite', 'corps', 'retour_corp', 'oral'])
pp_df = pd.DataFrame(columns=['identite', 'metiers', 'justif'])
bilan_df = pd.DataFrame(columns=['identite', 'score', 'integration', 'fatigue', 'endurance', 'stress', 'entretien', 'parcours', 'implication', 'objectif', 'bilan', 'conclusion'])

sites = {
    'Lille' : [feedback_df.copy(), bilan_df.copy(), pp_df.copy(), transport_df.copy(), logement_df.copy(), capacity_df.copy()],
    'Roubaix' : [feedback_df.copy(), bilan_df.copy(), pp_df.copy(), transport_df.copy(), logement_df.copy(), capacity_df.copy()],
    'Armentières' : [feedback_df.copy(), bilan_df.copy(), pp_df.copy(), transport_df.copy(), logement_df.copy(), capacity_df.copy()],
    'Saint-Omer' : [feedback_df.copy(), bilan_df.copy(), pp_df.copy(), transport_df.copy(), logement_df.copy(), capacity_df.copy()]
}


for file in day1.glob('*.xlsx'):
    raw_data = pd.read_excel(file)
    #print(raw_data.axes)

    for row in raw_data.iterrows():
        toAdd_feedback= pd.DataFrame({'identite':row[1]['Mon nom, mon prénom'], 'jour':'jour 1', 'score':row[1]['Mon ressenti'], 'marquant':row[1]['✏️Ce qui m\'a marqué, ce que j\'ai aimé ou pas aimé...'], 'activite':row[1]['Aujourd\'hui j\'ai participé à : Brisons la Glace, Le Contrat idéal, Dixit Express, La Charte du SAS\n\n💡Ce que je retiens de la journée, qu\'est-ce qui m\'a été utile ? Qu\'est-ce qui m\'a surpris ? Une...'],'satisfaction':row[1]['Affirmation 1'], 'besoin':row[1]['✏️ Une aide dont j\'aurai besoin à l\'E2C \n']}, index=[1])
        toAdd_transport= pd.DataFrame({'identite':row[1]['Mon nom, mon prénom'], 'transport':row[1]['🚌 Est-ce que je suis à l\'aise pour prendre les transports en commun ? \n\n'], 'duree':row[1]['🚲Combien de temps je mets pour venir à l\'E2C ?\n\n'], 'max':row[1]['🚗 Combien de temps de trajet je suis prêt(e) à faire pour me rendre en stage ou au travail ?']}, index=[1])
        toAdd_logement= pd.DataFrame({'identite':row[1]['Mon nom, mon prénom'], 'logement':row[1]['🏡 Où je vis ?'], 'stable':row[1]['🏢 Ma situation de logement est-elle stable ?'],'changement':row[1]['📦 Si je pouvais déménager dans l\'immédiat, est-ce que je le ferais ?'], 'justif':row[1]['📦Si oui, pourquoi ?'], 'enfant':row[1]['👶 Est-ce que j\'ai des enfants et est-ce que j\'ai un moyen de garde pour pouvoir être présent(e) à toutes les séances de l\'E2C ?']}, index=[1])
        sites[row[1]['Nom du site']][0] = pd.concat([sites[row[1]['Nom du site']][0], toAdd_feedback], ignore_index=True)
        sites[row[1]['Nom du site']][3] = pd.concat([sites[row[1]['Nom du site']][3], toAdd_transport], ignore_index=True)
        sites[row[1]['Nom du site']][4] = pd.concat([sites[row[1]['Nom du site']][4], toAdd_logement], ignore_index=True)

for file in day2.glob('*.xlsx'):
    raw_data = pd.read_excel(file)
    #print(raw_data.axes)
    for row in raw_data.iterrows():
        toAdd_feedback= pd.DataFrame({'identite':row[1]['Mon nom et prénom'], 'jour':'jour 2', 'score':row[1]['Mon ressenti'],'marquant':row[1]['✏️Ce qui m\'a marqué, ce que j\'ai aimé ou pas aimé...'], 'activite':row[1]['✏️ Qu\'est-ce que je retiens de la journée ? Qu\'est-ce qui m\'a été utile ? Qu\'est-ce qui m\'a surpris ? Une chose que j\'ai découverte sur moi aujourd\'hui ?'], 'satisfaction':row[1]['Affirmation 1'], 'objectif':row[1]['✏️ Ce que j\'ai retenu de l\'E2C, quel est l\'objectif du parcours ? Qu\'est-ce que je dois faire, à quoi je m\'engage ? Quelles vont être mes obligations en tant que stagiaire ?\n']}, index=[1])
        toAdd_capacity = pd.DataFrame({'identite':row[1]['Mon nom et prénom'],'corps':row[1]['🏃L\'activité corporelle du matin (marche / yoga / défi tour)\n'], 'retour_corp':row[1]['✏️ Qu\'est-ce que j\'ai retenu de l\'activité ? Comment elle s\'est passée ?'], 'oral':row[1]['Affirmation 12']}, index=[1])
        sites[row[1]['Nom du site']][0] = pd.concat([sites[row[1]['Nom du site']][0], toAdd_feedback], ignore_index=True)
        sites[row[1]['Nom du site']][5] = pd.concat([sites[row[1]['Nom du site']][5], toAdd_capacity], ignore_index=True)    #print(lille_feedback_df, lille_capacity_df)

for file in day3.glob('*.xlsx'):
    raw_data = pd.read_excel(file)
    #print(raw_data.axes)
    for row in raw_data.iterrows():
        toAdd_feedback = pd.DataFrame(
            {'identite': row[1]['Mon nom et prénom'], 'jour': 'jour 3', 'score': row[1]['Mon ressenti'],
             'marquant': row[1]['✏️Ce qui m\'a marqué, ce que j\'ai aimé ou pas aimé...'], 'activite': row[1][
                '✏️Qu\'est-ce que j\'ai retenu  ? Qu\'est-ce qui m\'a été utile ? Qu\'est-ce qui m\'a surpris ?'],
             'satisfaction': row[1]['Affirmation 1'], 'atout': row[1]['✏️ Une force que j\'ai découverte chez moi et quelque chose que je veux améliorer ?'] },
            index=[1])
        toAdd_pp = pd.DataFrame({'identite': row[1]['Mon nom et prénom'], 'metiers': row[1]['💼Le World Café des métiers — ma curiosité\n\nParmi les secteurs découverts, coche ceux qui t\'ont intéressé(e) :'], 'justif': row[1]['✏️ Ce qui m\'a le plus attiré dans un métier et pourquoi ?']}, index=[1])
        toAdd_bilan = pd.DataFrame({'identite': row[1]['Mon nom et prénom'], 'score': row[1]['Mon ressenti2'], 'integration': row[1]['✏️ Est-ce que j\'ai parlé aux autres, est-ce que je me suis senti(e) accepté(e) ?'], 'fatigue': row[1]['Mon ressenti3'], 'endurance': row[1]['✏️ Est-ce que je pense pouvoir tenir ce rythme sur un parcours long ?'], 'stress': row[1]['Mon ressenti4'], 'entretien': row[1]['Comment je me prépare à l\'entretien ? Est-ce que ça m\'inquiète ? Si oui, qu\'est-ce qui m\'inquiète ?'], 'parcours': row[1]['✏️ Comment j\'envisage mon parcours à l\'E2C si j\'y suis accepté(e) ?'], 'implication': row[1]['✏️ Comment je vais m\'impliquer ? Qu\'est-ce que je suis prêt(e) à donner ?'], 'objectif': row[1]['✏️ Mon objectif principal à la sortie du parcours E2C'], 'bilan': row[1]['🏁Mon bilan de ces 3 jours en quelques mots\n\nCe que j\'emporte de ce SAS : une découverte sur moi-même, une rencontre, une idée, une envie…'], 'conclusion': row[1]['Un dernier mot à ajouter ?']}, index=[1])
        sites[row[1]['Nom du site']][0] = pd.concat([sites[row[1]['Nom du site']][0], toAdd_feedback], ignore_index=True)
        sites[row[1]['Nom du site']][2] = pd.concat([sites[row[1]['Nom du site']][2], toAdd_pp], ignore_index=True)
        sites[row[1]['Nom du site']][1] = pd.concat([sites[row[1]['Nom du site']][1], toAdd_bilan], ignore_index=True)

for site, data in sites.items():
    filename = f'./output/sas_{site}.xlsx'
    with pd.ExcelWriter(filename, engine='openpyxl') as writer:
        data[0].to_excel(writer, sheet_name='bilan quotidien', header=['Stagiaire', 'Jour', 'Avis général sur la journée', 'Point marquant de la journée', 'Retour sur les activités',  'Satisfaction', 'Besoin exprimé (jour 1)', 'Objectif identifié (jour2)', 'Force mobilisable (jour3)'],index=False)

        ws1 = writer.book['bilan quotidien']

        style_table(ws1)
        auto_width(ws1)
        ws1.auto_filter.ref = ws1.dimensions
        ws1.freeze_panes = "A2"

        data[1].to_excel(writer, sheet_name='bilan générale', header=['Stagiaire', 'Avis général', 'Intégration dans le groupe', 'Niveau de fatigue', 'Capacité à suivre un parcours long', 'Niveau de stress', 'préparation à l\'entretien', 'Projection dans le parcours', 'Engagement personnel', 'Objectif à l\'issue du parcours E2C' ,'bilan du SAS', 'Dernier mot'], index=False)

        ws2 = writer.book['bilan générale']

        style_table(ws2)
        auto_width(ws2)
        ws2.auto_filter.ref = ws2.dimensions
        ws2.freeze_panes = "A2"

        data[2].to_excel(writer, sheet_name='projet professionnel', header=['Stagiaire', 'Metier(s) choisi(s)', 'Raisons du choix'],index=False)

        ws3 = writer.book['projet professionnel']

        style_table(ws3)
        auto_width(ws3)

        data[3].to_excel(writer, sheet_name='infos spéciciques', header=['Stagiaire', 'A l\'aise avec les transport en commun','Temps de trajet', 'Maximum acceptable'], startcol=0, index=False)
        data[4].to_excel(writer, sheet_name='infos spéciciques', header=['Stagiaire', 'Situation de logement', 'Stabilité', 'Envie de déménager', 'Raison de changement', 'Enfants et solution de garde'], startcol=5, index=False)
        data[5].to_excel(writer, sheet_name='infos spéciciques', header=['Stagiaire', 'Implication dans activité physique', 'Analyse de l\'activité physique', 'Aisance à \'oral' ], startcol=12, index=False)

        ws4 = writer.book['infos spéciciques']

        # appliquer style global
        style_table(ws4)
        auto_width(ws4)

